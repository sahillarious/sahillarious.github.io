"""Persist the kept postings to JSON and Excel."""
from __future__ import annotations

import json
import os
from datetime import datetime
from typing import Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .models import EXPORT_COLUMNS, JobPosting


def _ensure_dir(path: str) -> None:
    os.makedirs(path, exist_ok=True)


def _write_with_fallback(path: str, writer: Callable[[str], None]) -> str:
    """Write via `writer(path)`; if the file is locked (e.g. open in Excel),
    retry once with a timestamped name so results are never lost."""
    try:
        writer(path)
        return path
    except PermissionError:
        root, ext = os.path.splitext(path)
        alt = f"{root}_{datetime.now():%Y%m%d_%H%M%S}{ext}"
        print(f"  [storage] {os.path.basename(path)} is locked (open?) "
              f"-> writing {os.path.basename(alt)} instead")
        writer(alt)
        return alt


def save_json(
    jobs: Iterable[JobPosting], output_dir: str, filename: str = config.JSON_FILENAME
) -> str:
    _ensure_dir(output_dir)
    path = os.path.join(output_dir, filename)
    records = [j.to_record() for j in jobs]

    def _write(target: str) -> None:
        with open(target, "w", encoding="utf-8") as fh:
            json.dump(records, fh, indent=2, ensure_ascii=False)

    return _write_with_fallback(path, _write)


def save_excel(
    jobs: Iterable[JobPosting],
    output_dir: str,
    filename: str = config.EXCEL_FILENAME,
    columns: list[str] = EXPORT_COLUMNS,
) -> str:
    _ensure_dir(output_dir)
    path = os.path.join(output_dir, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for job in jobs:
        record = job.to_record()
        row = [record.get(col, "") for col in columns]
        ws.append(row)

    # Make the URL column clickable.
    if "url" in columns:
        url_idx = columns.index("url") + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=url_idx)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="2563EB", underline="single")

    _autosize(ws, columns)
    ws.freeze_panes = "A2"
    return _write_with_fallback(path, wb.save)


def _autosize(ws, columns: list[str], max_width: int = 60) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        longest = len(col_name)
        for cell in ws[letter][1:]:
            if cell.value:
                longest = max(longest, min(len(str(cell.value)), max_width))
        ws.column_dimensions[letter].width = min(longest + 2, max_width)


def save_all(jobs: list[JobPosting], output_dir: str) -> tuple[str, str]:
    return save_json(jobs, output_dir), save_excel(jobs, output_dir)


# Dropped jobs lead with WHY they were cut, so the reason is the first column.
_DROPPED_COLUMNS = ["drop_reason"] + EXPORT_COLUMNS


def save_dropped(jobs: list[JobPosting], output_dir: str) -> tuple[str, str]:
    """Write the filtered-out postings (with drop_reason) for auditing."""
    json_path = save_json(jobs, output_dir, filename=config.DROPPED_JSON_FILENAME)
    xlsx_path = save_excel(
        jobs, output_dir, filename=config.DROPPED_EXCEL_FILENAME,
        columns=_DROPPED_COLUMNS,
    )
    return json_path, xlsx_path
